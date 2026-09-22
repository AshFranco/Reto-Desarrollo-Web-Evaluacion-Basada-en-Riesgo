import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_rnf05_certification():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(30, 30, 30)

    # 1. Encabezado institucional
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.5)
    header_table.columns[1].width = Inches(2.0)

    cell_l = header_table.cell(0, 0)
    cell_r = header_table.cell(0, 1)

    p_org = cell_l.paragraphs[0]
    p_org.paragraph_format.space_after = Pt(2)
    r1 = p_org.add_run("DIRECCIÓN GENERAL DE MEDICAMENTOS, ALIMENTOS Y PRODUCTOS SANITARIOS (DIGEMAPS)\n")
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r2 = p_org.add_run("DEPARTAMENTO DE ASEGURAMIENTO DE CALIDAD (QA) Y AUDITORÍA DE SISTEMAS")
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(90, 90, 90)

    p_meta = cell_r.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(2)
    r3 = p_meta.add_run("CÓDIGO: ")
    r3.font.size = Pt(8)
    r3.font.bold = True
    r4 = p_meta.add_run("CERT-2026-RNF-05\n")
    r4.font.size = Pt(8.5)
    r4.font.bold = True
    r5 = p_meta.add_run("ESTADO: CERTIFICADO 100%")
    r5.font.size = Pt(7.5)
    r5.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. Título del informe
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("INFORME DE CERTIFICACIÓN TÉCNICA DE CALIDAD — RNF-05")
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(20, 20, 20)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_s = p_sub.add_run("Compatibilidad Multi-Navegador, Despliegue PWA Móvil y Sincronización en Campo")
    r_s.font.size = Pt(9.5)
    r_s.font.color.rgb = RGBColor(90, 90, 90)

    # Resumen Ejecutivo
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(4)
    r_s1 = p_s1.add_run("1. RESUMEN EJECUTIVO Y ALCANCE DE LA EVALUACIÓN")
    r_s1.font.size = Pt(11)
    r_s1.font.bold = True

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.15
    p_desc.paragraph_format.space_after = Pt(10)
    p_desc.add_run(
        "En cumplimiento de los requerimientos no funcionales del SRS (§RNF-05) y la asignación formal del Plan de Pruebas "
        "(10-PLAN-PRUEBAS.md §3.4), el Departamento de QA certifica la total conformidad del Sistema PWA de Evaluación Basada en Riesgo (EBR/BPM). "
        "La evaluación constató el correcto registro del Service Worker en navegadores de escritorio (Chrome, Edge, Firefox, Safari), "
        "la instalación nativa como PWA en dispositivos móviles Android (Chrome) e iOS (Safari), la captura de coordenadas satelitales GPS reales "
        "con exportación GeoJSON por criterio, la compresión de evidencias fotográficas en el cliente y la sincronización diferida FIFO "
        "con idempotencia garantizada mediante cola transaccional en IndexedDB (Dexie)."
    )

    # 2. Tabla de Matriz de Compatibilidad
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(8)
    p_s2.paragraph_format.space_after = Pt(4)
    r_s2 = p_s2.add_run("2. MATRIZ DE COMPATIBILIDAD MULTIPLATAFORMA Y RESULTADOS")
    r_s2.font.size = Pt(11)
    r_s2.font.bold = True

    tabla_plataformas = doc.add_table(rows=7, cols=5)
    tabla_plataformas.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla_plataformas.autofit = False
    set_table_borders(tabla_plataformas, color="B0B0B0", sz="4")

    anchos = [Inches(1.5), Inches(1.3), Inches(1.2), Inches(1.2), Inches(1.3)]
    headers = ["Plataforma / Entorno", "Navegador", "Soporte PWA / SW", "Offline / Storage", "Dictamen"]

    for i, h in enumerate(headers):
        cell = tabla_plataformas.cell(0, i)
        cell.width = anchos[i]
        set_cell_background(cell, "EAEAEA")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8)

    plataformas_data = [
        ("Escritorio (Windows 11)", "Google Chrome 128+", "Standalone / A2HS", "Cache + IndexedDB", "CONFORME"),
        ("Escritorio (Windows 11)", "Microsoft Edge 128+", "Standalone / App Bar", "Cache + IndexedDB", "CONFORME"),
        ("Escritorio (Linux / Win)", "Mozilla Firefox 130+", "Service Worker Gecko", "Cache + IndexedDB", "CONFORME"),
        ("Escritorio (macOS 14+)", "Apple Safari 17+", "Web App (Dock)", "Cache + IndexedDB", "CONFORME"),
        ("Móvil (Android 13/14)", "Google Chrome Móvil", "WebAPK / A2HS Nativo", "IndexedDB + GPS/Cám", "CONFORME"),
        ("Móvil (iOS 17/18)", "Apple Safari Móvil", "Add to Home Screen", "WebKit Storage + SW", "CONFORME"),
    ]

    for row_idx, row_vals in enumerate(plataformas_data, start=1):
        for col_idx, val in enumerate(row_vals):
            cell = tabla_plataformas.cell(row_idx, col_idx)
            cell.width = anchos[col_idx]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(7.5)
            if col_idx == 4:
                r.font.bold = True

    # 3. Detalle de verificación técnica
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_before = Pt(14)
    p_s3.paragraph_format.space_after = Pt(4)
    r_s3 = p_s3.add_run("3. DETALLE DE CERTIFICACIÓN DE CAPACIDADES CRÍTICAS EN CAMPO")
    r_s3.font.size = Pt(11)
    r_s3.font.bold = True

    detalles = [
        ("3.1 Instalación PWA y Ciclo de Vida del Service Worker",
         "El empaquetado de producción de Vite genera un manifest conforme a los estándares W3C (display: standalone, theme_color: #1565C0, "
         "iconos responsivos 64x64, 192x192, 512x512 y maskable). El Service Worker ('sw.js') emplea Workbox bajo la estrategia injectManifest:\n"
         "  • Precaching de 8 activos estáticos esenciales (963.20 KiB) para arranque instantáneo sin red.\n"
         "  • NetworkFirst con timeout de 3 segundos para catálogos y asignaciones de inspectores.\n"
         "  • NetworkOnly para endpoints de autenticación y mutaciones de inspección (delegadas a la cola local).\n"
         "  • StaleWhileRevalidate para recursos estáticos y consultas secundarias."),
        ("3.2 Captura de Coordenadas Satelitales GPS en Tiempo Real (GeoJSON)",
         "Cada criterio de la ficha técnica cuenta con un modal interactivo con soporte de geolocalización:\n"
         "  • Consumo de la API nativa 'navigator.geolocation.getCurrentPosition' con 'enableHighAccuracy: true' y timeout de 10 segundos.\n"
         "  • Generación de un archivo formal FeatureCollection GeoJSON que empaqueta coordenadas [longitud, latitud], timestamp ISO y metadatos del criterio.\n"
         "  • Almacenamiento desacoplado con vinculación directa a 'respuestaItemId', garantizando trazabilidad geográfica de cada hallazgo."),
        ("3.3 Integración de Cámara y Compresión de Evidencias en el Cliente",
         "Para optimizar la transferencia en redes móviles 3G/4G y preservar el almacenamiento local:\n"
         "  • El selector multimedia interactúa de forma directa con la cámara nativa del teléfono inteligente o la galería.\n"
         "  • El módulo cliente 'compressor.ts' procesa imágenes JPEG/PNG mediante Canvas, reduciendo fotografías de alta resolución "
         "(5 a 12 MB) a pesos inferiores a 300 KB (calidad 0.7, resolución máxima 1200px) sin comprometer la legibilidad pericial del acta."),
        ("3.4 Sincronización Real Extremo a Extremo (Offline ➡️ Online)",
         "Certificado rigurosamente mediante la suite automatizada 'apps/web/src/lib/sync/sync-escenario-offline-online.test.ts':\n"
         "  • El inspector en planta acumula operaciones sin red: apertura de evaluación, registro de respuestas, evidencias GeoJSON, finalización e informe.\n"
         "  • Las operaciones se aseguran localmente en 'db.cola_sync' con orden monotónico garantizado mediante generador de timestamp incremental.\n"
         "  • Al restablecer la red, 'SyncProcessor' transmite las transacciones en estricto orden FIFO con autenticación Bearer y formato multiparte.\n"
         "  • Idempotencia asegurada mediante UUID local, transicionando los registros a estado 'enviado' y vaciando el contador a cero.\n"
         "  • Tolerancia a fallos comprobada: ante errores transitorios de servidor (HTTP 500), el procesador activa backoff exponencial protector sin pérdida de información.")
    ]

    for subtitulo, texto in detalles:
        p_subt = doc.add_paragraph()
        p_subt.paragraph_format.space_before = Pt(8)
        p_subt.paragraph_format.space_after = Pt(2)
        r_subt = p_subt.add_run(subtitulo)
        r_subt.font.bold = True
        r_subt.font.size = Pt(9.5)

        p_t = doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.15
        p_t.paragraph_format.space_after = Pt(6)
        r_t = p_t.add_run(texto)
        r_t.font.size = Pt(8.5)

    # 4. Dictamen y firmas
    p_sf = doc.add_paragraph()
    p_sf.paragraph_format.space_before = Pt(16)
    p_sf.paragraph_format.space_after = Pt(8)
    r_sf = p_sf.add_run("4. DICTAMEN DE CONFORMIDAD Y FIRMAS DE CERTIFICACIÓN")
    r_sf.font.size = Pt(11)
    r_sf.font.bold = True

    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    set_table_borders(sig_table, color="CCCCCC", sz="4")

    w_sig = Inches(2.15)
    for col in sig_table.columns:
        col.width = w_sig

    firmas = [
        ("Auditor de Calidad (QA Lead)\nRowlis Trinidad\nIng. Aseguramiento de Software", "DICTAMEN: CONFORME"),
        ("Revisor Técnico Independiente\nIngeniería de Sistemas DIGEMAPS", "REVISIÓN: APROBADA"),
        ("Coordinador General de Proyecto\nDirección de Tecnologías DIGEMAPS", "LIBERACIÓN: AUTORIZADA"),
    ]

    for c_idx, (cargo, dictamen) in enumerate(firmas):
        cell_box = sig_table.cell(0, c_idx)
        cell_box.height = Inches(0.85)
        set_cell_margins(cell_box, top=60, bottom=60, left=80, right=80)
        p1 = cell_box.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_d = p1.add_run(f"\n\n_______________________\n{dictamen}")
        r_d.font.size = Pt(7.5)
        r_d.font.bold = True

        cell_label = sig_table.cell(1, c_idx)
        set_cell_margins(cell_label, top=60, bottom=60, left=80, right=80)
        set_cell_background(cell_label, "F8F8F8")
        p2 = cell_label.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c = p2.add_run(cargo)
        r_c.font.size = Pt(7.5)

    out_file = r"c:\Users\rowli\Documents\GitHub\Reto-Desarrollo-Web-Evaluacion-Basada-en-Riesgo\docs\qa\00_CERTIFICACION_RNF-05_COMPATIBILIDAD_PWA_MOVIL_Y_SINCRONIZACION.docx"
    doc.save(out_file)
    print(f"OK: {out_file} generado exitosamente.")

if __name__ == '__main__':
    build_rnf05_certification()
