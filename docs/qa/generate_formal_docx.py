import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Aplica bordes estándar, limpios y discretos a toda la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_formal_bug_report():
    doc = docx.Document()

    # Configuración de márgenes estándar (1 pulgada / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Estilo base
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(30, 30, 30)

    # =========================================================================
    # 1. ENCABEZADO INSTITUCIONAL FORMAL (Tabla de 2 columnas)
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.5)
    header_table.columns[1].width = Inches(2.0)

    cell_l = header_table.cell(0, 0)
    cell_r = header_table.cell(0, 1)

    p_org = cell_l.paragraphs[0]
    p_org.paragraph_format.space_after = Pt(2)
    r1 = p_org.add_run("DIRECCIÓN GENERAL DE MEDICAMENTOS, ALIMENTOS Y PRODUCTOS SANITARIOS\n")
    r1.font.size = Pt(8.5)
    r1.font.bold = True
    r2 = p_org.add_run("DEPARTAMENTO DE ASEGURAMIENTO DE CALIDAD (QA)")
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(90, 90, 90)

    p_meta = cell_r.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(2)
    r3 = p_meta.add_run("CÓDIGO: ")
    r3.font.size = Pt(8)
    r3.font.bold = True
    r4 = p_meta.add_run("DEF-2026-010\n")
    r4.font.size = Pt(8.5)
    r4.font.bold = True
    r5 = p_meta.add_run("VERSIÓN: 1.0 | ESTADO: CERRADO")
    r5.font.size = Pt(7.5)
    r5.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # 2. TÍTULO DEL DOCUMENTO
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("INFORME DE NO CONFORMIDAD / DEFECTO TÉCNICO")
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(20, 20, 20)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_s = p_sub.add_run("Sistema: Evaluación Basada en Riesgo (EBR / BPM) | Módulo: Administración de Usuarios")
    r_s.font.size = Pt(9.5)
    r_s.font.color.rgb = RGBColor(90, 90, 90)

    # =========================================================================
    # 3. FICHA TÉCNICA DEL DEFECTO (Tabla Formal)
    # =========================================================================
    meta_table = doc.add_table(rows=5, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="D1D5DB", sz="4")

    col_widths = [Inches(1.5), Inches(2.0), Inches(1.4), Inches(1.6)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    meta_data = [
        ("Identificador:", "DEF-2026-010", "Fecha de Reporte:", "13/09/2026"),
        ("Módulo:", "Gestión de Usuarios y Roles", "Fecha de Cierre:", "13/09/2026"),
        ("Severidad:", "Alta (Integridad de Acceso)", "Prioridad:", "P1 (Urgente)"),
        ("Tipo de Defecto:", "Funcional / Regla de Autorización", "Ambiente:", "Staging / QA Local"),
        ("Reportado Por:", "Equipo de Pruebas de Calidad", "Dictamen Final:", "RESUELTO / VERIFICADO"),
    ]

    for r_idx, (c0_t, c1_t, c2_t, c3_t) in enumerate(meta_data):
        cell0 = meta_table.cell(r_idx, 0)
        cell1 = meta_table.cell(r_idx, 1)
        cell2 = meta_table.cell(r_idx, 2)
        cell3 = meta_table.cell(r_idx, 3)

        set_cell_background(cell0, "F9FAFB")
        set_cell_background(cell2, "F9FAFB")

        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r = p0.add_run(c0_t)
        r.font.bold = True
        r.font.size = Pt(8.5)

        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r = p1.add_run(c1_t)
        r.font.size = Pt(8.5)

        p2 = cell2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(c2_t)
        r.font.bold = True
        r.font.size = Pt(8.5)

        p3 = cell3.paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r = p3.add_run(c3_t)
        r.font.size = Pt(8.5)
        if "RESUELTO" in c3_t:
            r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Función auxiliar para títulos de sección formales
    def add_section_title(number_and_name):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(number_and_name)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(20, 20, 20)

    # =========================================================================
    # 1. DESCRIPCIÓN DEL DEFECTO
    # =========================================================================
    add_section_title("1. DESCRIPCIÓN DEL DEFECTO")
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.15
    p_desc.paragraph_format.space_after = Pt(6)
    r = p_desc.add_run(
        "En el Panel Administrativo (/admin), dentro de la tabla del directorio de usuarios, el Administrador del Sistema "
        "disponía de las acciones operativas 'Cambiar rol' y 'Desactivar' habilitadas para su propio registro de cuenta. "
        "El sistema no implementaba una regla de restricción de auto-democión, permitiendo al usuario administrador modificar "
        "su rol a 'Administrador de Empresa'. Como consecuencia directa, el usuario perdía los privilegios de acceso "
        "administrativo inmediatamente tras la confirmación, quedando imposibilitado para gestionar el sistema desde la interfaz gráfica."
    )
    r.font.size = Pt(9.5)

    # =========================================================================
    # 2. PASOS DE REPRODUCCIÓN
    # =========================================================================
    add_section_title("2. PASOS PARA LA REPRODUCCIÓN DEL HALLAZGO")
    steps = [
        "Ingresar al sistema con la cuenta de Administrador (admin@digemaps.gob.do).",
        "Acceder al módulo administrativo (/admin) y seleccionar la pestaña 'Gestión de Usuarios y Roles'.",
        "Ubicar la fila correspondiente al usuario 'Administrador del Sistema'.",
        "Hacer clic sobre la acción 'Cambiar rol'.",
        "En la ventana emergente, seleccionar el rol 'Administrador de Empresa' y confirmar la acción.",
        "Comprobar que la solicitud responde con estado 200 OK y el usuario es degradado, perdiendo el acceso a la administración.",
    ]
    for idx, st in enumerate(steps, 1):
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Inches(0.25)
        p_st.paragraph_format.space_after = Pt(2)
        r_num = p_st.add_run(f"2.{idx} ")
        r_num.font.bold = True
        r_num.font.size = Pt(9)
        r_txt = p_st.add_run(st)
        r_txt.font.size = Pt(9)

    # =========================================================================
    # 3. COMPORTAMIENTO OBSERVADO VS. ESPERADO (Tabla Formal)
    # =========================================================================
    add_section_title("3. COMPORTAMIENTO OBSERVADO Y COMPORTAMIENTO ESPERADO")

    table_comp = doc.add_table(rows=2, cols=2)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_comp.autofit = False
    set_table_borders(table_comp, color="D1D5DB", sz="4")

    table_comp.columns[0].width = Inches(3.25)
    table_comp.columns[1].width = Inches(3.25)

    c_h1 = table_comp.cell(0, 0)
    c_h2 = table_comp.cell(0, 1)
    set_cell_background(c_h1, "F3F4F6")
    set_cell_background(c_h2, "F3F4F6")
    set_cell_margins(c_h1, 60, 60, 80, 80)
    set_cell_margins(c_h2, 60, 60, 80, 80)

    p = c_h1.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Comportamiento Observado (Fallo)")
    r.font.bold = True
    r.font.size = Pt(8.5)

    p = c_h2.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Comportamiento Requerido (Correcto)")
    r.font.bold = True
    r.font.size = Pt(8.5)

    c_b1 = table_comp.cell(1, 0)
    c_b2 = table_comp.cell(1, 1)
    set_cell_margins(c_b1, 80, 80, 80, 80)
    set_cell_margins(c_b2, 80, 80, 80, 80)

    p = c_b1.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "- Acciones operativas activas en la fila del administrador.\n"
        "- Endpoint PATCH /api/v1/usuarios/:id/rol permitía cambiar el rol a cualquier usuario sin verificar privilegios mínimos del sistema.\n"
        "- Pérdida irreversible de acceso de administración."
    )
    r.font.size = Pt(8.5)

    p = c_b2.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(
        "- Botones 'Cambiar rol' y 'Desactivar' deshabilitados para cuentas con rol ADMINISTRADOR.\n"
        "- Distintivo visual 'Administrador (Protegido)' con indicación de bloqueo.\n"
        "- La API debe retornar error 400 Bad Request si se solicita la revocación del rol del Administrador del Sistema."
    )
    r.font.size = Pt(8.5)

    # =========================================================================
    # 4. ANÁLISIS DE CAUSA RAÍZ
    # =========================================================================
    add_section_title("4. ANÁLISIS TÉCNICO DE CAUSA RAÍZ (RCA)")
    p_rca = doc.add_paragraph()
    p_rca.paragraph_format.line_spacing = 1.15
    p_rca.paragraph_format.space_after = Pt(6)
    r = p_rca.add_run(
        "El defecto se originó por la ausencia de una regla de inmutabilidad de la cuenta administrativa tanto en el "
        "servicio de backend como en los componentes de interfaz de usuario. En el backend (usuarios.service.ts), la función "
        "actualizarRol ejecutaba una transacción directa sobre la tabla usuario_rol sin validar previamente si el usuario "
        "poseía el rol ADMINISTRADOR. En el frontend (DashboardAdmin.tsx), la tabla renderizaba los botones de acción sin "
        "comprobar si el usuario en sesión o el registro correspondían al administrador, careciendo de la guarda defensiva necesaria."
    )
    r.font.size = Pt(9.5)

    # =========================================================================
    # 5. ACCIÓN CORRECTIVA IMPLEMENTADA
    # =========================================================================
    add_section_title("5. ACCIÓN CORRECTIVA Y CAMBIOS EN EL CÓDIGO")
    actions = [
        ("Backend (usuarios.service.ts):", 
         "Se incorporó una comprobación previa en los métodos actualizarRol y actualizarEstado. Si el usuario objetivo posee el rol ADMINISTRADOR, la operación se interrumpe lanzando una excepción BadRequestException('No se puede revocar el rol del Administrador del Sistema')."),
        ("Frontend (DashboardAdmin.tsx):", 
         "Se añadió la constante evaluadora esAdmin. Cuando es verdadera, se reemplaza el chip de rol por 'Administrador (Protegido)' y se deshabilitan los botones 'Cambiar rol' y 'Desactivar' junto con un mensaje explicativo en tooltip."),
        ("Base de Datos (seed-demo.ts):", 
         "Se configuró la asignación explícita e idempotente del rol ADMINISTRADOR para la cuenta admin@digemaps.gob.do durante la ejecución de semillas de prueba."),
    ]
    for label, desc in actions:
        p_act = doc.add_paragraph()
        p_act.paragraph_format.left_indent = Inches(0.25)
        p_act.paragraph_format.space_after = Pt(3)
        r_lbl = p_act.add_run(label + " ")
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9)
        r_dsc = p_act.add_run(desc)
        r_dsc.font.size = Pt(9)

    # =========================================================================
    # 6. VERIFICACIÓN Y REGISTRO DE PRUEBAS
    # =========================================================================
    add_section_title("6. VERIFICACIÓN Y RESULTADOS DE PRUEBA")

    table_verif = doc.add_table(rows=4, cols=3)
    table_verif.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_verif.autofit = False
    set_table_borders(table_verif, color="D1D5DB", sz="4")

    table_verif.columns[0].width = Inches(3.0)
    table_verif.columns[1].width = Inches(2.2)
    table_verif.columns[2].width = Inches(1.3)

    vh0 = table_verif.cell(0, 0)
    vh1 = table_verif.cell(0, 1)
    vh2 = table_verif.cell(0, 2)
    for vh, t in [(vh0, "Prueba Ejecutada"), (vh1, "Comando de Verificación"), (vh2, "Resultado")]:
        set_cell_background(vh, "F3F4F6")
        set_cell_margins(vh, 60, 60, 80, 80)
        p = vh.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        r.font.bold = True
        r.font.size = Pt(8.5)

    verif_items = [
        ("Pruebas Unitarias del Backend (API)", "npm.cmd test -w apps/api", "Aprobado (56/56)"),
        ("Pruebas Unitarias del Frontend (Web)", "npm.cmd test -- --run -w apps/web", "Aprobado (160/160)"),
        ("Compilación de Producción (Vite/TS)", "npm.cmd run build -w apps/web", "Exitoso (0 errores)"),
    ]
    for r_i, (t1, t2, t3) in enumerate(verif_items, 1):
        c0 = table_verif.cell(r_i, 0)
        c1 = table_verif.cell(r_i, 1)
        c2 = table_verif.cell(r_i, 2)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)
        set_cell_margins(c2, 60, 60, 80, 80)

        p = c0.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t1)
        r.font.size = Pt(8.5)

        p = c1.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t2)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(90, 90, 90)

        p = c2.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t3)
        r.font.bold = True
        r.font.size = Pt(8.5)

    # =========================================================================
    # 7. FIRMAS DE CONFORMIDAD Y APROBACIÓN (Tabla Formal de Firmas)
    # =========================================================================
    add_section_title("7. APROBACIÓN Y CONFORMIDAD")

    table_sign = doc.add_table(rows=2, cols=3)
    table_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_sign.autofit = False
    set_table_borders(table_sign, color="D1D5DB", sz="4")

    for row in table_sign.rows:
        for cell in row.cells:
            cell.width = Inches(2.16)
            set_cell_margins(cell, 80, 80, 80, 80)

    # Header firmas
    for idx, rol in enumerate(["Reportado por", "Corregido por", "Aprobado por (QA Lead)"]):
        cell = table_sign.cell(0, idx)
        set_cell_background(cell, "F9FAFB")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(rol)
        r.font.bold = True
        r.font.size = Pt(8)

    # Body firmas
    firmas = [
        ("Equipo de Pruebas\nFirma: __________________\nFecha: 13/09/2026"),
        ("Desarrollo Backend/Frontend\nFirma: __________________\nFecha: 13/09/2026"),
        ("Aseguramiento de Calidad\nFirma: __________________\nFecha: 13/09/2026"),
    ]
    for idx, f in enumerate(firmas):
        cell = table_sign.cell(1, idx)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(f)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(60, 60, 60)

    output_path = r"c:\Users\rowli\Documents\GitHub\Reto-Desarrollo-Web-Evaluacion-Basada-en-Riesgo\docs\qa\INFORME_DEFECTO_FORMAL_DEF-010.docx"
    doc.save(output_path)
    print(f"Documento formal guardado exitosamente en: {output_path}")

if __name__ == "__main__":
    build_formal_bug_report()
