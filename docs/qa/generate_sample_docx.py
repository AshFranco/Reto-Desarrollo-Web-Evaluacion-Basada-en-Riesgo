import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_bug_report():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
    
    # Styles
    PRIMARY_HEX = "2A6DB0"
    DARK_BLUE_HEX = "1D4E80"
    BG_LIGHT_HEX = "F0F4F8"
    GRAY_TEXT_HEX = "5B6572"
    BORDER_HEX = "CBD5E1"
    
    COLOR_PRIMARY = RGBColor(42, 109, 176)
    COLOR_DARK = RGBColor(29, 78, 128)
    COLOR_TEXT = RGBColor(26, 32, 39)
    COLOR_MUTED = RGBColor(91, 101, 114)
    COLOR_SUCCESS = RGBColor(46, 125, 50)
    
    # --- HEADER / BANNER ---
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(4.8)
    header_table.columns[1].width = Inches(2.0)
    
    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    
    set_cell_background(cell_left, "FFFFFF")
    set_cell_background(cell_right, "FFFFFF")
    
    p_org = cell_left.paragraphs[0]
    p_org.paragraph_format.space_after = Pt(2)
    r_org = p_org.add_run("MINISTERIO DE SALUD PÚBLICA — DIGEMAPS")
    r_org.font.size = Pt(8.5)
    r_org.font.bold = True
    r_org.font.color.rgb = COLOR_MUTED
    
    p_sys = cell_left.add_paragraph()
    p_sys.paragraph_format.space_after = Pt(0)
    r_sys = p_sys.add_run("Sistema de Evaluación Basada en Riesgo (EBR / BPM)")
    r_sys.font.size = Pt(11)
    r_sys.font.bold = True
    r_sys.font.color.rgb = COLOR_DARK
    
    p_doc = cell_right.paragraphs[0]
    p_doc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_doc.paragraph_format.space_after = Pt(2)
    r_doc = p_doc.add_run("INFORME DE DEFECTO TÉCNICO")
    r_doc.font.size = Pt(8)
    r_doc.font.bold = True
    r_doc.font.color.rgb = COLOR_PRIMARY
    
    p_code = cell_right.add_paragraph()
    p_code.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_code.paragraph_format.space_after = Pt(0)
    r_code = p_code.add_run("REGISTRO: DEF-010")
    r_code.font.size = Pt(10)
    r_code.font.bold = True
    r_code.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # --- TITLE ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("DEF-010: Falta de Blindaje y Pérdida de Privilegios por Auto-Modificación del Rol de Administrador")
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_DARK

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Reporte de aseguramiento de calidad (QA), análisis de causa raíz y certificación de resolución técnica.")
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_MUTED

    # --- METADATA TABLE ---
    meta_table = doc.add_table(rows=4, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    col_widths = [Inches(1.5), Inches(2.0), Inches(1.5), Inches(1.8)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            
    meta_fields = [
        ("Módulo Afectado:", "Gestión de Usuarios y Roles", "Severidad:", "Alta (Pérdida de privilegios)"),
        ("Tipo de Defecto:", "Funcional / Seguridad Operativa", "Prioridad:", "P1 (Urgente / Bloqueante)"),
        ("Ambiente:", "Staging / QA Local (Win 11 / Chrome)", "Estado:", "CERRADO / RESUELTO"),
        ("Reportado Por:", "Equipo de Pruebas de Usuario", "Fecha Cierre:", "13/09/2026"),
    ]
    
    for row_idx, (l1, v1, l2, v2) in enumerate(meta_fields):
        # Cell 0
        c0 = meta_table.cell(row_idx, 0)
        set_cell_background(c0, BG_LIGHT_HEX)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(l1)
        r0.font.bold = True
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = COLOR_DARK
        
        # Cell 1
        c1 = meta_table.cell(row_idx, 1)
        set_cell_background(c1, "FFFFFF")
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v1)
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = COLOR_TEXT
        
        # Cell 2
        c2 = meta_table.cell(row_idx, 2)
        set_cell_background(c2, BG_LIGHT_HEX)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(l2)
        r2.font.bold = True
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = COLOR_DARK
        
        # Cell 3
        c3 = meta_table.cell(row_idx, 3)
        set_cell_background(c3, "FFFFFF")
        p3 = c3.paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(v2)
        r3.font.size = Pt(8.5)
        if "CERRADO" in v2:
            r3.font.bold = True
            r3.font.color.rgb = COLOR_SUCCESS
        else:
            r3.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # Helper for section titles
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = COLOR_PRIMARY
        
        # Subtle horizontal divider under title
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_after = Pt(6)
        r_div = p_div.add_run("—" * 58)
        r_div.font.size = Pt(6)
        r_div.font.color.rgb = RGBColor(203, 213, 225)

    # --- 1. DESCRIPCIÓN ---
    add_section_header("1. Resumen y Contexto del Defecto")
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.paragraph_format.line_spacing = 1.15
    r_desc = p_desc.add_run(
        "En el Panel de Control Administrativo (/admin), dentro de la pestaña 'Gestión de Usuarios y Roles', "
        "el usuario que ostenta el rol de Administrador del Sistema figuraba en la tabla general con las acciones "
        "'Cambiar rol' y 'Desactivar' activas y operativas en su propia fila. Al hacer clic en 'Cambiar rol', "
        "el sistema permitía reasignarse un rol inferior (por ejemplo, 'Administrador de Empresa'), perdiendo "
        "de inmediato el acceso administrativo sin confirmación de seguridad ni bloqueo contra auto-democión."
    )
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = COLOR_TEXT

    # --- 2. PASOS PARA REPRODUCIR ---
    add_section_header("2. Pasos para Reproducir (Steps to Reproduce)")
    pasos = [
        "Iniciar sesión en la aplicación con las credenciales de administración (admin@digemaps.gob.do).",
        "Navegar al panel administrativo (/admin) y seleccionar la pestaña 'Gestión de Usuarios y Roles'.",
        "Localizar la fila correspondiente al usuario 'Administrador del Sistema'.",
        "Pulsar el botón de acción 'Cambiar rol' ubicado en dicha fila.",
        "Seleccionar en el desplegable modal el rol 'Administrador de Empresa' y confirmar.",
        "Comprobar que la solicitud finaliza exitosamente con código HTTP 200 OK y el administrador queda relegado sin posibilidad de recuperar sus privilegios desde la interfaz."
    ]
    for i, paso in enumerate(pasos, 1):
        p_paso = doc.add_paragraph()
        p_paso.paragraph_format.left_indent = Inches(0.2)
        p_paso.paragraph_format.space_after = Pt(3)
        r_num = p_paso.add_run(f"Paso {i}: ")
        r_num.font.bold = True
        r_num.font.size = Pt(9)
        r_num.font.color.rgb = COLOR_DARK
        r_txt = p_paso.add_run(paso)
        r_txt.font.size = Pt(9)
        r_txt.font.color.rgb = COLOR_TEXT

    # --- 3. COMPARATIVA: OBSERVADO VS ESPERADO ---
    add_section_header("3. Comportamiento Observado vs. Comportamiento Esperado")
    comp_table = doc.add_table(rows=2, cols=2)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False
    comp_table.columns[0].width = Inches(3.4)
    comp_table.columns[1].width = Inches(3.4)
    
    # Headers
    h1 = comp_table.cell(0, 0)
    h2 = comp_table.cell(0, 1)
    set_cell_background(h1, "FEE2E2") # Light red
    set_cell_background(h2, "DCFCE7") # Light green
    set_cell_margins(h1, 70, 70, 100, 100)
    set_cell_margins(h2, 70, 70, 100, 100)
    
    p_h1 = h1.paragraphs[0]
    p_h1.paragraph_format.space_after = Pt(0)
    r_h1 = p_h1.add_run("COMPORTAMIENTO OBSERVADO (DEFECTO)")
    r_h1.font.bold = True
    r_h1.font.size = Pt(8.5)
    r_h1.font.color.rgb = RGBColor(185, 28, 28)
    
    p_h2 = h2.paragraphs[0]
    p_h2.paragraph_format.space_after = Pt(0)
    r_h2 = p_h2.add_run("COMPORTAMIENTO ESPERADO (CORRECTO)")
    r_h2.font.bold = True
    r_h2.font.size = Pt(8.5)
    r_h2.font.color.rgb = RGBColor(21, 128, 61)
    
    # Body
    b1 = comp_table.cell(1, 0)
    b2 = comp_table.cell(1, 1)
    set_cell_margins(b1, 90, 90, 100, 100)
    set_cell_margins(b2, 90, 90, 100, 100)
    
    p_b1 = b1.paragraphs[0]
    p_b1.paragraph_format.space_after = Pt(0)
    p_b1.paragraph_format.line_spacing = 1.15
    r_b1 = p_b1.add_run(
        "• El botón 'Cambiar rol' y 'Desactivar' permanecían habilitados para la cuenta del Administrador.\n"
        "• La API (PATCH /usuarios/:id/rol) no validaba si el sujeto objeto de la mutación era el Administrador principal.\n"
        "• Auto-democión permitida: El administrador quedaba degradado a Empresa, perdiendo acceso total al panel."
    )
    r_b1.font.size = Pt(8.5)
    r_b1.font.color.rgb = COLOR_TEXT
    
    p_b2 = b2.paragraphs[0]
    p_b2.paragraph_format.space_after = Pt(0)
    p_b2.paragraph_format.line_spacing = 1.15
    r_b2 = p_b2.add_run(
        "• La fila del Administrador debe presentar una insignia 'Administrador (Protegido)' con ícono de candado.\n"
        "• Los botones 'Cambiar rol' y 'Desactivar' deben estar inhabilitados con tooltip explicativo de protección.\n"
        "• El Backend debe rechazar con HTTP 400 Bad Request cualquier intento de mutación contra la cuenta de administración."
    )
    r_b2.font.size = Pt(8.5)
    r_b2.font.color.rgb = COLOR_TEXT

    # --- 4. ANÁLISIS DE CAUSA RAÍZ ---
    add_section_header("4. Análisis Técnico de Causa Raíz (Root Cause Analysis)")
    p_rca = doc.add_paragraph()
    p_rca.paragraph_format.space_after = Pt(6)
    p_rca.paragraph_format.line_spacing = 1.15
    r_rca = p_rca.add_run(
        "La causa raíz radicó en la ausencia de una regla de autorización a nivel de dominio y presentación "
        "(Anti Self-Demotion / Admin Inmutability Rule). En el controlador y servicio de usuarios (usuarios.service.ts), "
        "el método actualizarRol ejecutaba directamente una transacción Prisma eliminando todos los roles del usuario "
        "e insertando el nuevo, sin consultar los roles preexistentes del idUsuario en cuestión. De forma análoga, "
        "en el componente React (DashboardAdmin.tsx), la tabla mapeaba todos los registros devueltos por el hook "
        "useUsuariosTodos sin aplicar un filtro de deshabilitación condicional basado en u.roles o u.correoElectronico."
    )
    r_rca.font.size = Pt(9.5)
    r_rca.font.color.rgb = COLOR_TEXT

    # --- 5. SOLUCIÓN IMPLEMENTADA ---
    add_section_header("5. Solución Técnica Implementada")
    soluciones = [
        ("Protección Backend (API NestJS):", 
         "En apps/api/src/modules/usuarios/usuarios.service.ts, se adicionó una verificación previa en actualizarRol y actualizarEstado. Si el usuario evaluado posee el rol ADMINISTRADOR, se bloquea la transacción lanzando BadRequestException('No se puede revocar el rol del Administrador del Sistema')."),
        ("Protección Frontend (UI React / MUI):", 
         "En apps/web/src/pages/dashboard/DashboardAdmin.tsx, se introdujo el detector esAdmin = u.roles.some(r => r.codigo === 'ADMINISTRADOR'). En la columna de rol se renderiza un Chip delineado 'Administrador (Protegido)' con LockOutlinedIcon y los botones de acción se renderizan como disabled envueltos en Tooltip protector."),
        ("Restauración en Semilla de Datos:", 
         "En apps/api/prisma/seed-demo.ts, se garantizó la eliminación y reasignación idempotente del rol ADMINISTRADOR para admin@digemaps.gob.do ante cada ejecución de prueba."),
    ]
    for tit, desc in soluciones:
        p_sol = doc.add_paragraph()
        p_sol.paragraph_format.space_after = Pt(4)
        p_sol.paragraph_format.left_indent = Inches(0.2)
        r_sol_t = p_sol.add_run(f"✔ {tit} ")
        r_sol_t.font.bold = True
        r_sol_t.font.size = Pt(9)
        r_sol_t.font.color.rgb = COLOR_DARK
        r_sol_d = p_sol.add_run(desc)
        r_sol_d.font.size = Pt(9)
        r_sol_d.font.color.rgb = COLOR_TEXT

    # --- 6. EVIDENCIA DE VERIFICACIÓN ---
    add_section_header("6. Evidencia de Pruebas y Dictamen de Cierre")
    verif_table = doc.add_table(rows=4, cols=3)
    verif_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    verif_table.autofit = False
    verif_table.columns[0].width = Inches(3.2)
    verif_table.columns[1].width = Inches(2.2)
    verif_table.columns[2].width = Inches(1.4)
    
    # Headers
    vh0 = verif_table.cell(0, 0)
    vh1 = verif_table.cell(0, 1)
    vh2 = verif_table.cell(0, 2)
    for vh, txt in [(vh0, "PRUEBA AUTOMATIZADA / VERIFICACIÓN"), (vh1, "COMANDO EJECUTADO"), (vh2, "RESULTADO")]:
        set_cell_background(vh, BG_LIGHT_HEX)
        set_cell_margins(vh, 70, 70, 100, 100)
        p = vh.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txt)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = COLOR_DARK
        
    filas_verif = [
        ("Suite de Pruebas Backend (Controladores y Servicios)", "npm.cmd test -w apps/api", "PASS (56/56)"),
        ("Suite de Pruebas Frontend (DashboardAdmin & Componentes)", "npm.cmd test -- --run -w apps/web", "PASS (160/160)"),
        ("Compilación de Producción TypeScript / Vite", "npm.cmd run build -w apps/web", "BUILD OK (0 err)"),
    ]
    for row_idx, (f1, f2, f3) in enumerate(filas_verif, 1):
        c0 = verif_table.cell(row_idx, 0)
        c1 = verif_table.cell(row_idx, 1)
        c2 = verif_table.cell(row_idx, 2)
        set_cell_margins(c0, 70, 70, 100, 100)
        set_cell_margins(c1, 70, 70, 100, 100)
        set_cell_margins(c2, 70, 70, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(f1)
        r0.font.size = Pt(8.5)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(f2)
        r1.font.size = Pt(8)
        r1.font.color.rgb = COLOR_MUTED
        
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(f3)
        r2.font.bold = True
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = COLOR_SUCCESS

    # Final verdict
    p_verdict = doc.add_paragraph()
    p_verdict.paragraph_format.space_before = Pt(12)
    p_verdict.paragraph_format.space_after = Pt(0)
    r_v1 = p_verdict.add_run("Dictamen Final de QA: ")
    r_v1.font.bold = True
    r_v1.font.size = Pt(10)
    r_v1.font.color.rgb = COLOR_DARK
    r_v2 = p_verdict.add_run("DEFECTO RESUELTO Y CERRADO SATISFACTORIAMENTE. Se certifica la no regresión y la integridad operativa del sistema.")
    r_v2.font.size = Pt(10)
    r_v2.font.bold = True
    r_v2.font.color.rgb = COLOR_SUCCESS

    output_path = r"c:\Users\rowli\Documents\GitHub\Reto-Desarrollo-Web-Evaluacion-Basada-en-Riesgo\docs\qa\EJEMPLO_INFORME_BUG_DEF-010.docx"
    doc.save(output_path)
    print(f"Documento guardado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_bug_report()
