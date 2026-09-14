import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_master_index():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(9.5)
    style.font.color.rgb = RGBColor(30, 30, 30)

    # Header
    ht = doc.add_table(rows=1, cols=2)
    ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    ht.columns[0].width = Inches(4.7)
    ht.columns[1].width = Inches(2.0)
    
    cl = ht.cell(0, 0)
    cr = ht.cell(0, 1)
    
    p = cl.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("DIRECCIÓN GENERAL DE MEDICAMENTOS, ALIMENTOS Y PRODUCTOS SANITARIOS\n")
    r.font.size = Pt(8.5)
    r.font.bold = True
    r2 = p.add_run("DEPARTAMENTO DE ASEGURAMIENTO DE CALIDAD (QA)")
    r2.font.size = Pt(8)
    r2.font.color.rgb = RGBColor(90, 90, 90)
    
    p2 = cr.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_after = Pt(2)
    r3 = p2.add_run("ÍNDICE MAESTRO\n")
    r3.font.size = Pt(8.5)
    r3.font.bold = True
    r4 = p2.add_run("TOTAL: 14 DEFECTOS CERRADOS")
    r4.font.size = Pt(7.5)
    r4.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Title
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_after = Pt(4)
    rt = p_t.add_run("ÍNDICE MAESTRO DE INFORMES DE DEFECTOS Y RESOLUCIONES")
    rt.font.size = Pt(13)
    rt.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    rsub = p_sub.add_run("Consolidado oficial de las 14 no conformidades detectadas, analizadas y resueltas durante el ciclo de pruebas de QA.")
    rsub.font.size = Pt(9.5)
    rsub.font.color.rgb = RGBColor(90, 90, 90)

    # Resumen Ejecutivo
    p_res = doc.add_paragraph()
    p_res.paragraph_format.space_after = Pt(8)
    p_res.paragraph_format.line_spacing = 1.15
    p_res.add_run(
        "El presente índice relaciona los 14 informes técnicos individuales generados en formato Microsoft Word (.docx), "
        "los cuales documentan el diagnóstico de causa raíz (RCA), la solución implementada a nivel de código y base de datos, "
        "y las evidencias de verificación automatizada. Todos los defectos reportados se encuentran actualmente en estado "
        "CERRADO / VERIFICADO con 100% de pruebas unitarias y de integración aprobadas (56 en Backend, 160 en Frontend)."
    )

    # Tabla maestra
    items = [
        ("DEF-2026-001", "Técnico Evaluador", "Reapertura de evaluación finalizada para corrección técnica", "Media", "Cerrado"),
        ("DEF-2026-002", "Motor de Riesgo", "Fallo HTTP 400 por ausencia de subcategorías de alimento", "Alta", "Cerrado"),
        ("DEF-2026-003", "Técnico Evaluador", "Falta de visibilidad de antecedentes e historial de inspección", "Media", "Cerrado"),
        ("DEF-2026-004", "Coordinador", "Desvinculación y reasignación de técnicos en casos activos", "Alta", "Cerrado"),
        ("DEF-2026-005", "Coordinador / Histórico", "Inspección detallada de expedientes cerrados e histórico", "Media", "Cerrado"),
        ("DEF-2026-006", "Portal Empresa", "Validación estricta contra signos negativos en RNC y teléfono", "Media", "Cerrado"),
        ("DEF-2026-007", "Establecimientos / API", "Desbordamiento de entero de 32 bits (P2020 / Error 500)", "Crítica", "Cerrado"),
        ("DEF-2026-008", "Portal Empresa", "Navegación y retorno cómodo a la pantalla de Mi Empresa", "Baja", "Cerrado"),
        ("DEF-2026-009", "Técnico Evaluador", "Salida cómoda del modo de evaluación al panel del técnico", "Media", "Cerrado"),
        ("DEF-2026-010", "Administración", "Blindaje contra auto-democión del rol de Administrador", "Alta", "Cerrado"),
        ("DEF-2026-011", "Administración", "Humanización de etiqueta 'Pendiente de validación' y estética", "Media", "Cerrado"),
        ("DEF-2026-012", "Coordinador", "Bandeja de informes devueltos y acción Deshacer devolución", "Alta", "Cerrado"),
        ("DEF-2026-013", "Coordinador", "Retiro de tooltip técnico interno y unificación de acciones", "Baja", "Cerrado"),
        ("DEF-2026-014", "Solicitud BPM / Catálogo", "Catálogo dinámico y desplegable de Tipos de Establecimiento", "Media", "Cerrado"),
    ]

    t = doc.add_table(rows=len(items) + 1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t, color="D1D5DB", sz="4")

    widths = [Inches(1.2), Inches(1.3), Inches(2.7), Inches(0.8), Inches(0.7)]
    for row in t.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[idx], 60, 60, 80, 80)

    # Headers
    h_titles = ["Código", "Módulo", "Descripción del Defecto", "Severidad", "Estado"]
    for idx, ht_title in enumerate(h_titles):
        c = t.cell(0, idx)
        set_cell_background(c, "F3F4F6")
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ht_title)
        r.font.bold = True
        r.font.size = Pt(8.5)

    for r_idx, (cod, mod, desc, sev, est) in enumerate(items, 1):
        c0 = t.cell(r_idx, 0)
        c1 = t.cell(r_idx, 1)
        c2 = t.cell(r_idx, 2)
        c3 = t.cell(r_idx, 3)
        c4 = t.cell(r_idx, 4)

        for cell, val, bold in [(c0, cod, True), (c1, mod, False), (c2, desc, False), (c3, sev, False), (c4, est, True)]:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(val)
            r.font.bold = bold
            r.font.size = Pt(8)

    # Sign-off
    p_sign_title = doc.add_paragraph()
    p_sign_title.paragraph_format.space_before = Pt(14)
    p_sign_title.paragraph_format.space_after = Pt(4)
    r = p_sign_title.add_run("APROBACIÓN Y CERTIFICACIÓN DEL PAQUETE DE CALIDAD")
    r.font.bold = True
    r.font.size = Pt(10)

    st = doc.add_table(rows=2, cols=3)
    st.alignment = WD_TABLE_ALIGNMENT.CENTER
    st.autofit = False
    set_table_borders(st, color="D1D5DB", sz="4")

    for row in st.rows:
        for cell in row.cells:
            cell.width = Inches(2.23)
            set_cell_margins(cell, 70, 70, 80, 80)

    for idx, h_txt in enumerate(["Preparado por", "Revisado por", "Aprobado por"]):
        cell = st.cell(0, idx)
        set_cell_background(cell, "F9FAFB")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_txt)
        r.font.bold = True
        r.font.size = Pt(8)

    signs = [
        ("Equipo de Aseguramiento QA\nFirma: __________________\nFecha: 13/09/2026"),
        ("Líder de Desarrollo de Software\nFirma: __________________\nFecha: 13/09/2026"),
        ("Dirección de Tecnología / DIGEMAPS\nFirma: __________________\nFecha: 13/09/2026"),
    ]
    for idx, s_txt in enumerate(signs):
        cell = st.cell(1, idx)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(s_txt)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(60, 60, 60)

    out_file = r"c:\Users\rowli\Documents\GitHub\Reto-Desarrollo-Web-Evaluacion-Basada-en-Riesgo\docs\qa\00_INDICE_MAESTRO_INFORMES_DEFECTOS.docx"
    doc.save(out_file)
    print("Indice maestro generado exitosamente!")

if __name__ == "__main__":
    build_master_index()
